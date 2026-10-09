"""Builds the pre-demo intro slide (16:9, Salesforce palette) for the Output Quality policy.

Usage: python3 build_intro_slide.py <output.pptx>
"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

BRAND = RGBColor(0x01, 0x76, 0xD3)
BRAND_DARK = RGBColor(0x01, 0x44, 0x86)
NAVY = RGBColor(0x03, 0x2D, 0x60)
CLOUD = RGBColor(0x1B, 0x96, 0xFF)
SKY = RGBColor(0x00, 0xA1, 0xE0)
TEXT = RGBColor(0x18, 0x18, 0x18)
TEXT_2 = RGBColor(0x44, 0x44, 0x44)
MUTED = RGBColor(0x74, 0x74, 0x74)
BORDER = RGBColor(0xC9, 0xC9, 0xC9)
BG = RGBColor(0xF3, 0xF3, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OK = RGBColor(0x2E, 0x84, 0x4A)
WARN = RGBColor(0xFE, 0x93, 0x39)
ERR = RGBColor(0xEA, 0x00, 0x1E)

FONT = "Salesforce Sans"


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def text(slide, x, y, w, h, runs, size=14, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None):
    """runs: str, or list of paragraphs; each paragraph is a str or a list of (text, overrides)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    paragraphs = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        parts = para if isinstance(para, list) else [(para, {})]
        for chunk, o in parts:
            r = p.add_run()
            r.text = chunk
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.color.rgb = o.get("color", color)
    return tb


def arrow(slide, x1, y1, x2, y2, color=BRAND, width=2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def card(slide, x, y, w, h, kicker, title):
    box(slide, x, y, w, h, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    text(slide, x + 0.25, y + 0.18, w - 0.5, 0.3, kicker, size=11, color=BRAND, bold=True)
    text(slide, x + 0.25, y + 0.44, w - 0.5, 0.4, title, size=17, color=BRAND_DARK, bold=True)


def node(slide, x, y, w, h, title, sub, fill=WHITE, line=BRAND, title_color=BRAND_DARK, sub_color=TEXT_2):
    box(slide, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    text(slide, x + 0.1, y + 0.06, w - 0.2, h - 0.12,
         [[(title, {"bold": True, "color": title_color, "size": 12})], [(sub, {"size": 10, "color": sub_color})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def status_row(slide, x, y, color, glyph, label, desc):
    circle = box(slide, x, y + 0.04, 0.42, 0.42, fill=color, shape=MSO_SHAPE.OVAL)
    tf = circle.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = glyph
    r.font.name = FONT
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = WHITE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    text(slide, x + 0.58, y - 0.02, 3.0, 0.3, label, size=14, bold=True, color=TEXT)
    text(slide, x + 0.58, y + 0.26, 3.0, 0.3, desc, size=11, color=TEXT_2)


def build(path):
    prs = Presentation()
    prs.slide_width = Emu(12192000)   # 13.333 in
    prs.slide_height = Emu(6858000)   # 7.5 in
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG

    # Title band
    box(s, 0, 0, 13.333, 1.35, fill=NAVY)
    box(s, 0, 1.35, 13.333, 0.05, fill=SKY)
    text(s, 0.6, 0.22, 12, 0.3, "OMNI GATEWAY POLICY  ·  FOR AI AGENTS", size=12, color=CLOUD, bold=True)
    text(s, 0.6, 0.5, 12, 0.6, "Output Quality & Anomaly Detection", size=32, color=WHITE, bold=True)

    text(s, 0.6, 1.62, 12.2, 0.45,
         "Catches hallucinated, biased, toxic and off-track AI answers on the gateway — before they reach users.",
         size=17, color=BRAND_DARK)

    top, h = 2.3, 3.9
    # 1. Where it runs
    card(s, 0.6, top, 3.9, h, "01  WHERE IT RUNS", "Outbound, on the gateway")
    node(s, 0.85, top + 1.05, 1.95, 0.55, "User / app", "A2A or LLM API call")
    node(s, 0.85, top + 1.95, 1.95, 0.7, "Omni Gateway", "outbound policy scores the reply",
         fill=NAVY, line=NAVY, title_color=WHITE, sub_color=CLOUD)
    node(s, 0.85, top + 3.0, 1.95, 0.6, "AI agent / LLM", "e.g. CloudHub 2.0 agent")
    arrow(s, 1.825, top + 1.6, 1.825, top + 1.95)
    arrow(s, 1.825, top + 2.65, 1.825, top + 3.0)
    node(s, 3.0, top + 1.95, 1.3, 0.7, "LLM judge", "OpenAI gpt-5.4", line=CLOUD)
    arrow(s, 2.8, top + 2.3, 3.0, top + 2.3, color=CLOUD, width=1.5)

    # 2. How it decides
    card(s, 4.72, top, 3.9, h, "02  HOW IT DECIDES", "Three layers, always on")
    steps = [
        ("1", "LLM judge", "Scores each answer 0–1 for hallucination, toxicity, bias and anomaly — and sees the user's question."),
        ("2", "Word-list fallback", "Scores the same risks if the judge is slow or down. Traffic is never held back."),
        ("3", "Signals & feedback", "Looping, empty or truncated output, failed tasks, and spikes in users' 👎."),
    ]
    for i, (n, t, d) in enumerate(steps):
        y = top + 1.0 + i * 0.95
        text(s, 4.97, y - 0.04, 0.4, 0.45, n, size=24, bold=True, color=BRAND)
        text(s, 5.4, y, 3.0, 0.3, t, size=14, bold=True, color=TEXT)
        text(s, 5.4, y + 0.28, 3.0, 0.62, d, size=11, color=TEXT_2)

    # 3. What it does
    card(s, 8.84, top, 3.9, h, "03  WHAT IT DOES", "Your choice, per category")
    status_row(s, 9.09, top + 1.05, OK, "✓", "Pass", "Clean answers flow untouched")
    status_row(s, 9.09, top + 1.8, WARN, "!", "Annotate", "Delivered with a quality report")
    status_row(s, 9.09, top + 2.55, ERR, "×", "Block", "Answer withheld from the user")
    text(s, 9.09, top + 3.3, 3.45, 0.5,
         "Every flag is logged by the gateway and reported in the response.", size=11, color=MUTED)

    # Stats strip
    sy = 6.38
    stats = [
        ("1.1–2.4 s", "LLM verdict per answer"),
        ("4", "risk categories"),
        ("A2A v0.3 & v1.0", "plus OpenAI & Anthropic APIs"),
        ("102", "automated tests"),
    ]
    sw = 12.14 / len(stats)
    for i, (big, small) in enumerate(stats):
        x = 0.6 + i * sw
        if i:
            box(s, x - 0.02, sy + 0.08, 0.02, 0.5, fill=BORDER)
        text(s, x + 0.1, sy, sw - 0.2, 0.4, big, size=20, bold=True, color=BRAND, align=PP_ALIGN.CENTER)
        text(s, x + 0.1, sy + 0.4, sw - 0.2, 0.3, small, size=11, color=TEXT_2, align=PP_ALIGN.CENTER)

    text(s, 0.6, 7.12, 12.2, 0.25,
         "Numbers from the live CloudHub 2.0 run on agent-network-ingress-gw, 9 Oct 2026 (judge: OpenAI gpt-5.4).",
         size=9, color=MUTED)
    prs.save(path)


if __name__ == "__main__":
    build(sys.argv[1])
